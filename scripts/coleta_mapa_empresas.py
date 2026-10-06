"""Baixa os XLSX mensais do Mapa de Empresas (MEMP/DREI) e soma por mes.

Fonte: https://www.gov.br/memp/pt-br/acesso-a-informacao/dados-abertos/repositorio-de-dados-abertos
Cada arquivo mensal traz 1 linha por municipio (aba "Abertas, Fechadas e
Ativas"); somando as colunas tem-se o total BR do mes de referencia.
O portal mantem so os meses recentes — rode 1x/mes; o historico acumula
em data/manual/manual_empresas.csv (consumido por coleta_conjuntura.py).

Uso:  python scripts/coleta_mapa_empresas.py
Saida: data/manual/manual_empresas.csv  (mes,abertas,fechadas; mes = dia 01)
"""
import re
import sys
import urllib.request
from datetime import date
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
MANUAL = BASE / "data" / "manual"
MANUAL.mkdir(exist_ok=True)

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) ObservatorioDosDados"}

REPO = ("https://www.gov.br/memp/pt-br/acesso-a-informacao/"
        "dados-abertos/repositorio-de-dados-abertos")
RE_XLSX = re.compile(r"([\w\-]+mapa-de-empresas[a-z0-9\-]*\.xlsx)/@@download/file")


def get_text(url: str) -> str:
    req = urllib.request.Request(url, headers=UA)
    return urllib.request.urlopen(req, timeout=120).read().decode("utf-8", "replace")


def soma_mes(raw: bytes):
    import openpyxl
    import io
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    ws = wb["Abertas, Fechadas e Ativas"]
    ref = None
    ab = fe = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # cabecalho
        if not row or row[0] is None:
            continue
        d = row[0]
        dstr = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d)[:10]
        ref = ref or dstr
        try:
            ab += int(row[4] or 0)
        except (ValueError, TypeError):
            pass
        try:
            fe += int(row[5] or 0)
        except (ValueError, TypeError):
            pass
    wb.close()
    if not ref:
        return None
    y, m, _d = ref.split("-")
    return f"{y}-{m}-01", ab, fe


def perfil_mes(raw: bytes):
    """Retrato do arquivo mais recente: ativas por NJ/porte, genero por UF,
    tempo medio de abertura (quando o portal voltar a publicar numeros)."""
    import openpyxl
    import io
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    ref = None
    nj_porte, genero = [], []
    tempos = []
    try:
        ws = wb["Empresas Ativas por NJ e Porte"]
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row or row[0] is None:
                continue
            d = row[0]
            ref = ref or (d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d)[:10])
            try:
                nj_porte.append({"nj": str(row[2] or "")[:60], "porte": str(row[3] or "")[:40],
                                 "ativas": int(row[4] or 0)})
            except (ValueError, TypeError):
                pass
    except KeyError:
        pass
    try:
        ws = wb["Gênero (UF, NJ e Porte)"]
    except KeyError:
        try:
            ws = wb[[t for t in wb.sheetnames if "nero" in t][0]]
        except IndexError:
            ws = None
    if ws is not None:
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row or row[2] is None:
                continue
            try:
                genero.append({"uf": str(row[3] or "")[:2], "genero": str(row[5] or "")[:20],
                               "ativas": int(row[6] or 0)})
            except (ValueError, TypeError):
                pass
    try:
        ws = wb["Abertas, Fechadas e Ativas"]
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row or len(row) < 8:
                continue
            try:
                tempos.append(float(row[7]))
            except (ValueError, TypeError):
                pass
    except KeyError:
        pass
    wb.close()
    out = {"ref": ref, "nj_porte": nj_porte, "genero_uf": genero}
    if tempos:
        out["tempo_medio_dias"] = round(sum(tempos) / len(tempos), 1)
    return out


if __name__ == "__main__":
    import json
    html = get_text(REPO)
    arqs = sorted(set(RE_XLSX.findall(html)))
    print(f"[mapa] {len(arqs)} arquivos mensais no portal.")
    achados = {}
    raws = {}
    for a in arqs:
        url = f"https://www.gov.br/memp/pt-br/acesso-a-informacao/dados-abertos/{a}/@@download/file"
        try:
            req = urllib.request.Request(url, headers=UA)
            raw = urllib.request.urlopen(req, timeout=180).read()
            raws[a] = raw
            r = soma_mes(raw)
            if r:
                achados[r[0]] = (r[1], r[2])
                print(f"[mapa] {r[0]}: abertas={r[1]:,} fechadas={r[2]:,}")
        except Exception as e:
            print(f"[mapa] FALHOU {a}: {str(e)[:100]}")
    # retrato (perfil) do arquivo mais recente
    if raws:
        novo = sorted(raws)[-1]
        perfil = perfil_mes(raws[novo])
        PUB = BASE / "public" / "data"
        PUB.mkdir(parents=True, exist_ok=True)
        (PUB / "conj_perfil.json").write_text(
            json.dumps(perfil, ensure_ascii=False), encoding="utf-8")
        print(f"[mapa] perfil {perfil['ref']}: {len(perfil['nj_porte'])} NJ/porte, "
              f"{len(perfil['genero_uf'])} linhas genero/UF"
              + (f", tempo medio {perfil['tempo_medio_dias']} dias"
                 if perfil.get("tempo_medio_dias") else " (tempo em revisao no portal)"))
    dest = MANUAL / "manual_empresas.csv"
    prev = {}
    if dest.exists():
        for ln in dest.read_text(encoding="utf-8").splitlines()[1:]:
            p = ln.split(",")
            if len(p) >= 3 and p[0].strip():
                try:
                    prev[p[0].strip()[:10]] = (int(float(p[1])), int(float(p[2])))
                except ValueError:
                    pass
    prev.update(achados)
    linhas = [f"{m},{v[0]},{v[1]}" for m, v in sorted(prev.items())]
    dest.write_text("mes,abertas,fechadas\n" + "\n".join(linhas) + "\n", encoding="utf-8")
    print(f"OK: {dest} com {len(linhas)} meses "
          f"({linhas[0][:7]} -> {linhas[-1][:7]}). Rode coleta_conjuntura.py p/ publicar.")
